"""Stateless composition and export for one retained workspace result page."""

from __future__ import annotations

import html

from kindred.core.quantity_units import (
    CUBIC_METRE,
    MOLAR,
    SECOND,
    Unit,
    analysis_display_unit,
    compatible_units,
    convert_value,
    format_unit_for_display,
    parse_unit,
)

import io
import csv
import json
import unicodedata
from dataclasses import dataclass, fields, is_dataclass, replace
from enum import Enum
from types import MappingProxyType
from typing import BinaryIO, Callable, Iterator, Mapping, Sequence
from zipfile import ZIP_DEFLATED, ZipFile

import numpy as np
from openpyxl import Workbook
from openpyxl.cell.cell import WriteOnlyCell
from openpyxl.writer.excel import ExcelWriter

from kindred.core.project_presentation_state import (
    AxisSeriesSelection,
    ExperimentalAxisSelection,
    VesselVolumeAxisSelection,
    SimulationSeriesAxisSelection,
)
from kindred.core.project_document import ProjectDocument
from kindred.core.result_repository import (
    ResultRepository,
    SimulationResultValue,
)
from kindred.core.simulator.solver_types import SimulationPresentationBoundary
from kindred.core.simulator.dsl_parse import PreparedDslRequestSet
from kindred.gui.curve_copy_headers import observable_copy_definitions
from kindred.gui.dataset_payload_snapshot import ExperimentalLayerInput
from kindred.gui.plot_config import AxisSelectionCandidate
from kindred.core.result_curve_projection import ProjectedCurveChunk, physical_curve_chunks

__all__ = [
    "InteractiveDisplayCurve",
    "InteractiveDisplayProjection",
    "ResultDataBlock",
    "ResultLayerKind",
    "ResultRenderInput",
    "VisibleResultSet",
    "XlsxExportCompletion",
    "compose_canonical_export_blocks",
    "compose_interactive_display",
    "compose_result_input",
    "encode_clipboard_text",
    "encode_tsv_row",
    "iter_blocks_tsv_chunks",
    "validate_xlsx_export",
    "write_xlsx",
]

_MAX_COPY_CELLS = 1_000_000
_MAX_COPY_CHARACTERS = 64 * 1024 * 1024
_MAX_XLSX_STRING_LENGTH = 32_767
_EXCEL_MAX_ROWS = 1_048_576
_EXCEL_MAX_COLUMNS = 16_384


def _readonly_display_array(value: object) -> np.ndarray:
    """Retain one admitted transient display array without copying it."""

    array = np.asarray(value)
    array.setflags(write=False)
    return array


@dataclass(frozen=True, slots=True)
class InteractiveDisplayCurve:
    """One immutable, transient physical curve; never Copy/export authority."""

    key: tuple[object, ...]
    label: str
    x: np.ndarray
    y: np.ndarray
    # Canonical seconds stay aligned when X is another quantity or is sampled.
    sample_time: np.ndarray
    color: str
    source: ResultDataBlock
    dashed: bool = False
    opacity: float = 1.0
    # These arrays are display topology, not result authority.  Keeping them
    # with the transient curve prevents the preview renderer from connecting
    # points across an intervention/discontinuity while it is cross-fading.
    segment_ids: np.ndarray | None = None
    segment_keys: tuple[str, ...] | None = None
    discontinuity_ids: np.ndarray | None = None
    physical_chunks: tuple[ProjectedCurveChunk, ...] = ()

    def __post_init__(self) -> None:
        x = _readonly_display_array(self.x)
        y = _readonly_display_array(self.y)
        segment_ids = (
            None
            if self.segment_ids is None
            else _readonly_display_array(self.segment_ids)
        )
        discontinuity_ids = (
            None
            if self.discontinuity_ids is None
            else _readonly_display_array(self.discontinuity_ids)
        )
        object.__setattr__(self, "x", x)
        object.__setattr__(self, "y", y)
        object.__setattr__(self, "sample_time", _readonly_display_array(self.sample_time))
        object.__setattr__(self, "segment_ids", segment_ids)
        object.__setattr__(self, "discontinuity_ids", discontinuity_ids)


_DisplayAnnotation = tuple[str, str, Mapping[str, object], float]


@dataclass(frozen=True, slots=True)
class InteractiveDisplayProjection:
    tab_id: str
    set_id: str
    x_label: str
    y_label: str
    curves: tuple[InteractiveDisplayCurve, ...]
    intervention_annotations: tuple[_DisplayAnnotation, ...] = ()

    def __post_init__(self) -> None:
        object.__setattr__(self, "curves", tuple(self.curves))
        object.__setattr__(self, "intervention_annotations", tuple(self.intervention_annotations))


class _XlsxDestination:
    def __init__(
        self,
        destination: BinaryIO,
        *,
        cancellation_check: Callable[[], None] | None,
    ) -> None:
        if not callable(getattr(destination, "write", None)):
            raise TypeError("XLSX destination must be a writable binary stream.")
        if cancellation_check is not None and not callable(cancellation_check):
            raise TypeError("cancellation_check must be callable.")
        self._destination = destination
        self._cancellation_check = cancellation_check
        self._maximum_size = 0

    def write(self, data: object) -> int:
        if self._cancellation_check is not None:
            self._cancellation_check()
        payload = bytes(data)
        written = int(self._destination.write(payload))
        if written != len(payload):
            raise OSError(
                f"XLSX destination accepted only {written} of {len(payload)} bytes."
            )
        self._maximum_size = max(self._maximum_size, int(self.tell()))
        return written

    def __getattr__(self, name: str) -> object:
        return getattr(self._destination, name)

    @property
    def bytes_written(self) -> int:
        return self._maximum_size


@dataclass(frozen=True, slots=True)
class XlsxExportCompletion:
    worksheet_count: int
    data_row_count: int
    bytes_written: int


def _cleanup_write_only_workbook(workbook: Workbook) -> list[BaseException]:
    """Close writer streams and remove every owned worksheet tempfile."""

    failures: list[BaseException] = []
    for sheet in workbook.worksheets:
        rows = getattr(sheet, "_rows", None)
        if rows is not None:
            try:
                rows.close()
            except (StopIteration, RuntimeError, ValueError):
                pass
            except BaseException as exc:
                failures.append(exc)
        writer = getattr(sheet, "_writer", None)
        if writer is None:
            continue
        try:
            writer.close()
        except (StopIteration, RuntimeError, ValueError):
            pass
        except BaseException as exc:
            failures.append(exc)
        try:
            writer.cleanup()
        except (FileNotFoundError, ValueError):
            pass
        except BaseException as exc:
            failures.append(exc)
    try:
        workbook.close()
    except BaseException as exc:
        failures.append(exc)
    return failures


def _required_id(value: object, field: str) -> str:
    normalized = str(value or "").strip()
    if not normalized:
        raise ValueError(f"{field} must be nonblank.")
    return normalized


def _axis_selection(value: object, *, field: str) -> AxisSeriesSelection:
    if not isinstance(
        value,
        (SimulationSeriesAxisSelection, ExperimentalAxisSelection, VesselVolumeAxisSelection),
    ):
        raise TypeError(
            f"{field} must be a typed simulation or experimental axis selection."
        )
    return value


def _is_time_selection(value: AxisSeriesSelection) -> bool:
    return isinstance(value, SimulationSeriesAxisSelection) and value.series_name == "t"


def _json_safe(value: object) -> object:
    if isinstance(value, Mapping):
        return {str(key): _json_safe(item) for key, item in value.items()}
    if isinstance(value, (tuple, list)):
        return [_json_safe(item) for item in value]
    if isinstance(value, np.ndarray):
        return [_json_safe(item) for item in value.tolist()]
    if isinstance(value, np.generic):
        return value.item()
    if isinstance(value, Enum):
        return _json_safe(value.value)
    if is_dataclass(value) and not isinstance(value, type):
        return {
            item.name: _json_safe(getattr(value, item.name)) for item in fields(value)
        }
    if value is None or isinstance(value, (bool, int, float, str)):
        return value
    return str(value)


def encode_tsv_row(values: Sequence[object]) -> str:
    """Encode one RFC-style quoted tab-separated row with a final newline."""

    output = io.StringIO(newline="")
    writer = csv.writer(
        output,
        dialect="excel-tab",
        lineterminator="\n",
        quoting=csv.QUOTE_ALL,
    )
    writer.writerow(tuple("" if value is None else str(value) for value in values))
    return output.getvalue()


def encode_clipboard_text(value: str) -> str:
    """Protect one dynamic clipboard text cell from spreadsheet formulas."""

    if not isinstance(value, str):
        raise TypeError("clipboard text must be a string.")
    index = 0
    while index < len(value):
        character = value[index]
        if character.isspace() or unicodedata.category(character).startswith("C"):
            index += 1
            continue
        if character in "=+-@":
            return "'" + value
        break
    return value


def _canonical_json(value: object) -> str:
    return json.dumps(
        _json_safe(value),
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    )


def _series_tuple(
    result: SimulationResultValue,
) -> tuple[tuple[str, np.ndarray], ...]:
    projected = [(name, values) for name, values in result.series.items()]
    existing_names = set(result.series)
    for name in result.succeeded_algebra_outputs:
        if name in existing_names or name not in result.scalars:
            continue
        projected.append(
            (
                name,
                _readonly_display_array(
                    np.full(result.t.shape, result.scalars[name])
                ),
            )
        )
    return tuple(projected)


def _series_names(result: SimulationResultValue) -> tuple[str, ...]:
    """Return projected display names without allocating scalar arrays."""

    names = list(result.series)
    existing_names = set(names)
    names.extend(
        name
        for name in result.succeeded_algebra_outputs
        if name in result.scalars and name not in existing_names
    )
    return tuple(names)


def _result_warning_details(result: SimulationResultValue) -> tuple[str, ...]:
    details: list[str] = []
    if result.failed_algebra_outputs:
        details.append(
            "Algebra failed: " + ", ".join(result.failed_algebra_outputs)
        )
    for cause in result.algebra_errors:
        raw_name = cause.details.get("name")
        observable_name = str(raw_name or "<unknown observable>")
        code = cause.code if cause.code is not None else "<none>"
        details.append(
            "Algebra error for "
            f"{observable_name}: type={cause.type}; code={code}; "
            f"message={cause.message}"
        )
    if result.omitted_algebra_outputs:
        details.append(
            "Algebra omitted: " + ", ".join(result.omitted_algebra_outputs)
        )
    return tuple(details)


def _result_provenance_payload(result: SimulationResultValue) -> Mapping[str, object]:
    """Serialize readable provenance directly from the canonical result."""

    return MappingProxyType(
        {
            "provenance": MappingProxyType(
                {
                    "solver": result.solver_provenance,
                    "termination": result.termination,
                    "volume": result.volume,
                    "boundaries": result.presentation_boundaries,
                    "algebra": MappingProxyType(
                        {
                            "requested": result.requested_algebra_outputs,
                            "succeeded": result.succeeded_algebra_outputs,
                            "failed": result.failed_algebra_outputs,
                            "omitted": result.omitted_algebra_outputs,
                            "errors": tuple(
                                cause.to_payload()
                                for cause in result.algebra_errors
                            ),
                        }
                    ),
                }
            ),
        }
    )


@dataclass(frozen=True, slots=True)
class VisibleResultSet:
    set_id: str
    label: str

    def __post_init__(self) -> None:
        object.__setattr__(self, "set_id", _required_id(self.set_id, "set_id"))
        object.__setattr__(self, "label", _required_id(self.label, "label"))


class ResultLayerKind(Enum):
    EXPLICIT = "explicit"
    PREVIEW = "preview"
    REFERENCE = "reference"
    EXPERIMENTAL = "experimental"


@dataclass(frozen=True, slots=True)
class ResultDataBlock:
    """One immutable full-resolution result/data table and its render style."""

    block_id: str | ExperimentalAxisSelection
    label: str
    kind: ResultLayerKind
    t: np.ndarray
    # Unlike plotted X (t), this remains time in seconds; None means unrecorded.
    sample_time: np.ndarray | None
    series: tuple[tuple[str, np.ndarray], ...]
    series_axes: tuple[tuple[str, AxisSeriesSelection], ...]
    x_axis: AxisSeriesSelection
    result: SimulationResultValue | None
    segment_ids: np.ndarray | None
    segment_keys: tuple[str, ...] | None
    discontinuity_ids: np.ndarray | None
    source_row: np.ndarray | None
    x_name: str = "t"
    time_column_name: str | None = None
    copy_definitions: tuple[tuple[str, str], ...] = ()
    quantity_units: tuple[tuple[str, Unit | None], ...] = ()
    presentation_units: tuple[tuple[str, Unit | None], ...] = ()
    series_provenance: tuple[tuple[str, Mapping[str, object]], ...] = ()
    x_provenance: Mapping[str, object] = MappingProxyType({})
    color_by_series: tuple[tuple[str, str], ...] = ()
    visible: bool = True
    dashed: bool = False
    opacity: float = 1.0
    markers: bool = False
    point_size: float = 6.0
    point_shape: str = "o"
    presentation_boundaries: tuple[SimulationPresentationBoundary, ...] = ()
    first_index: int = 0

    def unit_for(self, name: str, *, x: bool = False) -> Unit | None:
        if self.kind is ResultLayerKind.EXPERIMENTAL:
            return SECOND if x and isinstance(self.x_axis, SimulationSeriesAxisSelection) and self.x_axis.series_name == "t" else dict(self.quantity_units)[name]
        if x and isinstance(self.x_axis, VesselVolumeAxisSelection):
            return CUBIC_METRE
        if not x and isinstance(dict(self.series_axes).get(name), VesselVolumeAxisSelection):
            return CUBIC_METRE
        return dict(self.quantity_units).get(name)

    def unit_label(self, name: str, *, x: bool = False) -> str:
        unit = self.unit_for(name, x=x)
        if (
            x and self.kind is ResultLayerKind.EXPERIMENTAL
            and _is_time_selection(self.x_axis) and self.time_column_name is not None
        ):
            name = self.time_column_name
        return f"{name} [{unit.symbol}]" if unit is not None else name

    def __post_init__(self) -> None:
        # Blocks are built only by this module's composers; keep the genuine
        # cross-field invariants and immutability, not generic type checks.
        time = self.t
        names: set[str] = set()
        for name, values in self.series:
            if name in names:
                raise ValueError(f"data block contains duplicate series {name!r}.")
            if values.size != time.size and not (
                self.kind is ResultLayerKind.EXPERIMENTAL and time.size == 0
            ):
                raise ValueError(f"data series {name!r} length does not match time.")
            names.add(name)
        if not names:
            raise ValueError("data block requires at least one series.")
        series_axes = tuple(self.series_axes)
        if tuple(name for name, _ in series_axes) != tuple(name for name, _ in self.series):
            raise ValueError("data series axes must align exactly with data series.")
        if len({selection for _, selection in series_axes}) != len(series_axes):
            raise ValueError("data series axis identities must be unique.")
        if self.kind is ResultLayerKind.EXPERIMENTAL:
            if self.result is not None:
                raise ValueError("experimental data blocks prohibit repository entries.")
            if any(value is not None for value in (self.segment_ids, self.segment_keys, self.discontinuity_ids)):
                raise ValueError("experimental data blocks prohibit solver identity.")
            boundaries = ()
        else:
            if self.source_row is not None:
                raise ValueError("simulation data blocks prohibit observation source_row.")
            if self.segment_ids is None or self.segment_keys is None or self.discontinuity_ids is None:
                raise ValueError("simulation data blocks require complete executor identity.")
            boundaries = tuple(self.presentation_boundaries)
            if boundaries is not self.result.presentation_boundaries:
                raise ValueError("simulation data blocks must retain canonical presentation boundaries.")
        provenance_names = [name for name, _payload in self.series_provenance]
        if len(provenance_names) != len(set(provenance_names)) or not names.issuperset(provenance_names):
            raise ValueError("data series provenance must uniquely reference data.")
        colors = tuple(self.color_by_series)
        if not names.issubset({name for name, _color in colors}):
            raise ValueError("data series colors must cover every data series.")
        if len(colors) != len({name for name, _color in colors}):
            raise ValueError("data series colors must be unique by series.")
        if not np.isfinite(self.opacity) or not 0.0 <= self.opacity <= 1.0:
            raise ValueError("opacity must be between zero and one.")
        if not np.isfinite(self.point_size) or self.point_size <= 0.0:
            raise ValueError("point_size must be positive.")
        object.__setattr__(self, "series", tuple(self.series))
        object.__setattr__(self, "series_axes", series_axes)
        object.__setattr__(self, "series_provenance", tuple(
            (name, MappingProxyType({
                str(key): tuple(value) if isinstance(value, (list, tuple)) else value
                for key, value in payload.items()
            }))
            for name, payload in self.series_provenance
        ))
        object.__setattr__(self, "x_provenance", MappingProxyType({
            str(key): tuple(value) if isinstance(value, (list, tuple)) else value
            for key, value in self.x_provenance.items()
        }))
        object.__setattr__(self, "color_by_series", colors)
        object.__setattr__(self, "presentation_boundaries", boundaries)

def _select_analysis_block(
    blocks: Sequence[ResultDataBlock],
    current_set_id: str,
) -> ResultDataBlock | None:
    return next(
        (
            block
            for block in blocks
            if block.block_id == current_set_id
            and block.kind in {ResultLayerKind.EXPLICIT, ResultLayerKind.PREVIEW}
            and isinstance(block.result, SimulationResultValue)
        ),
        next(
            (
                block
                for block in blocks
                if block.block_id == current_set_id
                and block.kind is ResultLayerKind.REFERENCE
                and isinstance(block.result, SimulationResultValue)
            ),
            None,
        ),
    )


@dataclass(frozen=True, slots=True)
class ResultRenderInput:
    """Call-local render facts; never durable result or export authority."""

    tab_id: str
    current_set_id: str
    blocks: tuple[ResultDataBlock, ...]
    statistics_series_names: tuple[str, ...] = ()
    selected_x_series: AxisSeriesSelection = SimulationSeriesAxisSelection("t")
    selected_y_series: tuple[AxisSeriesSelection, ...] = ()
    mapped_overlay_locks: tuple[tuple[ExperimentalAxisSelection, str], ...] = ()
    selected_y_uses_default_all: bool = True
    x_label: str = "Time (s)"
    y_label: str = "Value"
    x_unit: Unit | None = None
    y_unit: Unit | None = None
    x_unit_options: tuple[Unit, ...] = ()
    y_unit_options: tuple[Unit, ...] = ()
    axis_candidates: tuple[AxisSelectionCandidate, ...] = ()
    reference_lines_visible: bool = True
    intervention_annotations: tuple[_DisplayAnnotation, ...] = ()
    intervention_annotations_visible: bool = True
    intervention_labels_visible: bool = True
    unavailable_x_diagnostic: str | None = None
    completion_warning_details: tuple[tuple[str, tuple[str, ...]], ...] = ()
    # Composer-resolved facts that interactive Preview composition reuses, so
    # it never re-derives membership, colours, or the simulation X axis.
    plot_set_labels: tuple[tuple[str, str], ...] = ()
    color_by_axis_selection: tuple[tuple[AxisSeriesSelection, str], ...] = ()
    simulation_x_series: AxisSeriesSelection = SimulationSeriesAxisSelection("t")

    def __post_init__(self) -> None:
        blocks = tuple(self.blocks)
        identities = tuple((block.kind, block.block_id, tuple(axis for _, axis in block.series_axes)) for block in blocks)
        if len(identities) != len(set(identities)):
            raise ValueError("render blocks must have unique semantic identities.")
        object.__setattr__(self, "blocks", blocks)
        names = tuple(self.statistics_series_names)
        if len(names) != len(set(names)):
            raise ValueError("statistics series names must be unique.")
        object.__setattr__(self, "statistics_series_names", names)
        selected_x = self.selected_x_series
        selected_y = tuple(self.selected_y_series)
        if len(selected_y) != len(set(selected_y)):
            raise ValueError("selected_y_series must be unique.")
        overlay_locks = tuple(self.mapped_overlay_locks)
        if len({axis for axis, _ in overlay_locks}) != len(overlay_locks):
            raise ValueError("mapped overlay locks must be unique by axis.")
        candidates = tuple(self.axis_candidates)
        candidate_ids = tuple(item.selection for item in candidates)
        if len(candidate_ids) != len(set(candidate_ids)) or selected_x not in candidate_ids:
            raise ValueError("axis candidates must be unique and include selected X.")
        if any(axis not in candidate_ids for axis, _ in overlay_locks):
            raise ValueError("mapped overlay locks must be axis candidates.")
        annotations = tuple(self.intervention_annotations)
        warnings = tuple(
            (label, tuple(details))
            for label, details in self.completion_warning_details
        )
        object.__setattr__(self, "selected_x_series", selected_x)
        object.__setattr__(self, "selected_y_series", selected_y)
        object.__setattr__(self, "mapped_overlay_locks", overlay_locks)
        object.__setattr__(self, "axis_candidates", candidates)
        object.__setattr__(self, "intervention_annotations", annotations)
        object.__setattr__(self, "completion_warning_details", warnings)

    @property
    def statistics_result(self) -> SimulationResultValue | None:
        block = _select_analysis_block(self.blocks, self.current_set_id)
        return block.result if block is not None else None

    @property
    def authored_volume_unit(self) -> Unit | None:
        result = self.statistics_result
        if result is None or result.volume is None:
            result = next((block.result for block in self.blocks
                           if block.result is not None and block.result.volume is not None), None)
        return None if result is None else dict(result.presentation_units)["Vessel volume"]

    @property
    def volume_unit_options(self) -> tuple[Unit, ...]:
        preferred = tuple(dict(block.result.presentation_units)["Vessel volume"]
                          for block in self.blocks if block.result is not None and block.result.volume is not None)
        return compatible_units(CUBIC_METRE, preferred) if preferred else ()

    def parameters_for_display(
        self, *, concentration_unit: Unit, time_unit: Unit, volume_unit: Unit,
    ) -> dict[str, tuple[float | None, str, str, str]]:
        """Project canonical parameters using the plot's Analysis unit choices."""
        result = self.statistics_result
        if result is None:
            return {}
        classifications = {
            name: (family, role)
            for name, family, role in result.parameter_classifications
        }
        rows = {}
        for name, value, unit_text in result.parameters:
            source = parse_unit(unit_text)
            display = analysis_display_unit(
                source, concentration=concentration_unit, time=time_unit, volume=volume_unit,
            )
            try:
                display_value = convert_value(value, source, display)
            except ValueError:
                display_value = None
            rows[name] = (display_value, display.symbol, *classifications[name])
        return rows

    @property
    def completion_warning_tooltip(self) -> str:
        return "\n\n".join(
            f"{label}:\n" + "\n".join(f"  • {detail}" for detail in details)
            for label, details in self.completion_warning_details
        )



def _result_block(
    *,
    set_id: str,
    label: str,
    result: SimulationResultValue,
    prepared_dsl: PreparedDslRequestSet,
    kind: ResultLayerKind,
    color_by_axis_selection: Mapping[AxisSeriesSelection, str],
    visible: bool = True,
) -> ResultDataBlock:
    display_label = (
        f"{label} [Preview]"
        if kind is ResultLayerKind.PREVIEW
        else f"{label} [Reference]"
        if kind is ResultLayerKind.REFERENCE
        else label
    )
    series = _series_tuple(result)
    series_axes = tuple(
        (name, SimulationSeriesAxisSelection(name)) for name, _values in series
    )
    return ResultDataBlock(
        block_id=set_id, label=display_label, kind=kind, t=result.t,
        sample_time=result.t,
        copy_definitions=observable_copy_definitions(prepared_dsl),
        quantity_units=result.quantity_units,
        presentation_units=result.presentation_units,
        series=series, series_axes=series_axes,
        x_axis=SimulationSeriesAxisSelection("t"), result=result,
        segment_ids=result.segment_ids, segment_keys=result.segment_keys,
        discontinuity_ids=result.discontinuity_ids, source_row=None,
        color_by_series=tuple(
            (selection.series_name, str(color))
            for selection, color in color_by_axis_selection.items()
            if isinstance(selection, SimulationSeriesAxisSelection)
        ),
        visible=visible, dashed=kind is ResultLayerKind.REFERENCE,
        opacity=0.42 if kind is ResultLayerKind.REFERENCE else 1.0,
        presentation_boundaries=result.presentation_boundaries,
    )



def _volume_block(block: ResultDataBlock) -> ResultDataBlock | None:
    result = block.result
    if result.volume is None:
        return None
    first = result.volume.first_index
    name = "Vessel volume"
    return replace(
        block, t=result.t[first:], sample_time=result.t[first:],
        series=((name, result.volume.values),),
        series_axes=((name, VesselVolumeAxisSelection()),),
        segment_ids=result.segment_ids[first:], segment_keys=result.segment_keys[first:], discontinuity_ids=result.discontinuity_ids[first:],
        first_index=first, color_by_series=((name, "#8b5cf6"),),
        quantity_units=(("t", SECOND), (name, CUBIC_METRE)),
        presentation_units=((name, dict(result.presentation_units)[name]),),
    )

def compose_canonical_export_blocks(
    *,
    document: ProjectDocument,
    tab_id: str,
    repository: ResultRepository,
) -> tuple[ResultDataBlock, ...]:
    """Project every retained active-tab result and its qualifying raw data."""

    if not isinstance(document, ProjectDocument):
        raise TypeError("Canonical export requires a ProjectDocument.")
    tab = document.tab(_required_id(tab_id, "tab_id"))
    if repository.tab_id != tab.tab_id:
        raise ValueError("Canonical export repository belongs to another tab.")

    blocks: list[ResultDataBlock] = []
    retained_set_ids: set[str] = set()
    preview_set_ids = set(repository.preview_set_ids())
    for initial_set in tab.initial_condition_sets:
        set_id = initial_set.set_id
        current = repository.current_result(set_id)
        reference = repository.reference_result(set_id)
        if current is not None:
            retained_set_ids.add(set_id)
            kind = (
                ResultLayerKind.PREVIEW
                if set_id in preview_set_ids
                else ResultLayerKind.EXPLICIT
            )
            blocks.append(
                _result_block(
                    set_id=set_id,
                    label=initial_set.name,
                    result=current,
                    prepared_dsl=repository.current_request(set_id).prepared_dsl,
                    kind=kind,
                    color_by_axis_selection={
                        SimulationSeriesAxisSelection(name): "#000000"
                        for name in _series_names(current)
                    },
                )
            )
        if reference is not None:
            retained_set_ids.add(set_id)
            blocks.append(
                _result_block(
                    set_id=set_id,
                    label=initial_set.name,
                    result=reference,
                    prepared_dsl=repository.reference_request(set_id).prepared_dsl,
                    kind=ResultLayerKind.REFERENCE,
                    color_by_axis_selection={
                        SimulationSeriesAxisSelection(name): "#000000"
                        for name in _series_names(reference)
                    },
                )
            )

    links = dict(tab.dataset_links)
    for record in document.datasets:
        link = links.get(record.dataset_id)
        if link is None or link.set_id not in retained_set_ids:
            continue
        mapping = dict(link.observation_mapping)
        for series in record.observations:
            source = ExperimentalLayerInput(
                axis=ExperimentalAxisSelection(
                    record.dataset_id,
                    series.raw_observation_id,
                ),
                dataset_name=record.display_name,
                raw_column=series.column_name,
                unit=series.unit,
                time_column_name=record.time_column_name,
                model_target=mapping.get(series.raw_observation_id),
                set_id=link.set_id,
                t=series.t,
                y=series.y,
                source_rows=series.source_rows,
            )
            label = f"{source.dataset_name}: {source.raw_column}"
            if source.model_target is not None:
                label = f"{label} → {source.model_target}"
            time = np.array(source.t, dtype=float, copy=True)
            values = np.array(source.y, dtype=float, copy=True)
            source_rows = np.array(source.source_rows, dtype=np.int64, copy=True)
            time.setflags(write=False)
            values.setflags(write=False)
            source_rows.setflags(write=False)
            blocks.append(
                ResultDataBlock(
                    block_id=source.axis,
                    label=label,
                    kind=ResultLayerKind.EXPERIMENTAL,
                    t=time,
                    sample_time=time if time.size else None,
                    time_column_name=source.time_column_name,
                    series=((source.raw_column, values),),
                    quantity_units=((source.raw_column, source.unit),),
                    series_axes=((source.raw_column, source.axis),),
                    x_axis=SimulationSeriesAxisSelection("t"),
                    result=None,
                    segment_ids=None,
                    segment_keys=None,
                    discontinuity_ids=None,
                    source_row=source_rows,
                    series_provenance=(
                        (source.raw_column, _observation_provenance(source)),
                    ),
                    color_by_series=((source.raw_column, "#000000"),),
                )
            )
    return tuple(blocks)


def _data_block(
    block: ResultDataBlock,
    *,
    x_axis: AxisSeriesSelection,
    y_axes: set[AxisSeriesSelection] | None = None,
) -> ResultDataBlock | None:
    from kindred.core.result_curve_projection import axis_values

    series_axis_map = dict(block.series_axes)
    if block.kind is ResultLayerKind.EXPERIMENTAL:
        return block if _is_time_selection(x_axis) else None
    coordinates = axis_values(block.result, x_axis)
    if coordinates is None:
        return None
    x_first, x_values, x_name = coordinates
    first = max(block.first_index, x_first)
    offset = first - block.first_index
    selected = tuple(
        (name, values[offset:]) for name, values in block.series
        if (series_axis_map[name] in y_axes if y_axes is not None
            else not isinstance(series_axis_map[name], VesselVolumeAxisSelection))
    )
    if not selected:
        return None
    names = {name for name, _ in selected}
    return replace(
        block, t=x_values[first-x_first:], sample_time=block.result.t[first:],
        series=selected,
        series_axes=tuple((name, selection) for name, selection in block.series_axes if name in names),
        x_axis=x_axis, x_name=x_name, first_index=first,
        segment_ids=block.result.segment_ids[first:], segment_keys=block.result.segment_keys[first:],
        discontinuity_ids=block.result.discontinuity_ids[first:],
        quantity_units=block.result.quantity_units,
        color_by_series=tuple((name, color) for name, color in block.color_by_series if name in names),
    )


def _has_selected_y_series(
    block: ResultDataBlock,
    *,
    y_axes: set[AxisSeriesSelection] | None,
) -> bool:
    return any(
        y_axes is None or selection in y_axes
        for _name, selection in block.series_axes
    )


def _statistics_series_names(
    block: ResultDataBlock,
    *,
    y_axes: set[AxisSeriesSelection] | None,
) -> tuple[str, ...]:
    succeeded = set(block.result.succeeded_algebra_outputs)
    return tuple(
        name for name in block.result.series
        if name not in succeeded
        and (
            y_axes is None
            or SimulationSeriesAxisSelection(name) in y_axes
        )
    )


def _selected_x_unavailable_diagnostic(
    selected_x: AxisSeriesSelection,
    unavailable_labels: Sequence[str],
) -> str | None:
    if _is_time_selection(selected_x) or not unavailable_labels:
        return None
    labels = ", ".join(repr(str(label)) for label in unavailable_labels)
    selected_identity = (
        f"series {selected_x.series_name!r}"
        if isinstance(selected_x, SimulationSeriesAxisSelection)
        else "vessel volume" if isinstance(selected_x, VesselVolumeAxisSelection)
        else f"raw observation {selected_x.raw_observation_id!r} in dataset {selected_x.dataset_id!r}"
    )
    return (
        f"Selected X {selected_identity} is unavailable for result layer(s) {labels}. "
        "Those layer(s) are omitted from the current plot only; "
        "full-resolution Copy All/XLSX exports remain unaffected."
    )


def _export_column_header(
    block: ResultDataBlock,
    name: str,
    provenance: Mapping[str, object],
    *, x: bool = False,
) -> str:
    header = f"{block.label}::{block.unit_label(name, x=x)}"
    identity = {
        key: str(provenance[key])
        for key in (
            "dataset_id",
            "raw_observation_id",
            "dataset_name",
            "raw_column",
            "model_target",
            "set_id",
        )
        if provenance.get(key) not in {None, ""}
    }
    if identity:
        encoded = json.dumps(identity, sort_keys=True, separators=(",", ":"))
        header = f"{header} [provenance={encoded}]"
    return encode_clipboard_text(header)


def _xlsx_metadata_rows(
    blocks: Sequence[ResultDataBlock],
) -> tuple[tuple[str, int, str], ...]:
    used: set[str] = set()
    seen: set[int] = set()
    rows: list[tuple[str, int, str]] = []
    for block in blocks:
        title = _worksheet_title(block.label, used)
        entry = block.result
        if entry is None or id(entry) in seen:
            continue
        seen.add(id(entry))
        serialized = _canonical_json(_result_provenance_payload(entry))
        rows.extend(
            (title, part, serialized[start:start + _MAX_XLSX_STRING_LENGTH])
            for part, start in enumerate(
                range(0, len(serialized), _MAX_XLSX_STRING_LENGTH), start=1
            )
        )
    return tuple(rows)


def _xlsx_literal_cell(sheet: object, value: object) -> object:
    """Single XLSX cell boundary: strings are always literal, never formulas."""

    if not isinstance(value, str):
        return value
    length = len(value)
    if length > _MAX_XLSX_STRING_LENGTH:
        raise ValueError(
            "XLSX string cell exceeds Excel's physical limit of "
            f"{_MAX_XLSX_STRING_LENGTH:,} characters (got {length:,})."
        )
    cell = WriteOnlyCell(sheet, value=value)
    cell.data_type = "s"
    return cell


def _append_xlsx_row(sheet: object, values: Sequence[object]) -> None:
    append = getattr(sheet, "append", None)
    if not callable(append):
        raise TypeError("XLSX worksheet must provide append().")
    append([_xlsx_literal_cell(sheet, value) for value in values])


def _validate_xlsx_worksheet_dimensions(
    dimensions: Sequence[tuple[str, int, int]],
) -> None:
    """Validate every planned worksheet against Excel's physical limits."""

    for title, rows, columns in dimensions:
        if rows > _EXCEL_MAX_ROWS:
            raise ValueError(
                f"XLSX worksheet {title!r} requires {rows:,} rows; "
                f"Excel permits at most {_EXCEL_MAX_ROWS:,}."
            )
        if columns > _EXCEL_MAX_COLUMNS:
            raise ValueError(
                f"XLSX worksheet {title!r} requires {columns:,} columns; "
                f"Excel permits at most {_EXCEL_MAX_COLUMNS:,}."
            )



def _boundary_export_columns(result):
    if result is None:
        return ()
    columns = []
    if result.volume is not None:
        first = result.volume.first_index
        columns.extend((
            ("Volume canonical row", np.arange(first, len(result.t)), "int"),
            ("Volume time [s]", result.t[first:], "float"),
            ("Vessel volume [m³]", result.volume.values, "float"),
        ))
    states = [(b.time, i, j, state) for i, b in enumerate(result.presentation_boundaries)
              for j, state in enumerate(b.states)]
    if states:
        columns.extend((
            ("Boundary time [s]", tuple(item[0] for item in states), "float"),
            ("Boundary index", tuple(item[1] for item in states), "int"),
            ("Boundary state ordinal", tuple(item[2] for item in states), "int"),
            ("Boundary vessel volume [m³]", tuple(item[3].volume for item in states), "float"),
            ("Boundary executed temperature [K]", tuple(item[3].temperature_K for item in states), "float"),
        ))
        columns.extend((f"Boundary {name} [M]", tuple(dict(item[3].concentrations)[name] for item in states), "float")
                       for name, _ in states[0][3].concentrations)
    return tuple(columns)

def _planned_xlsx_worksheet_dimensions(
    blocks: Sequence[ResultDataBlock],
    *,
    lineage_row_count: int,
    metadata_row_count: int,
) -> tuple[tuple[str, int, int], ...]:
    """Plan data and support worksheet dimensions without allocating them."""

    used: set[str] = set()
    dimensions: list[tuple[str, int, int]] = []
    for block in blocks:
        title = _worksheet_title(block.label, used)
        identity_column_count = 2 if block.kind is ResultLayerKind.EXPERIMENTAL else 4
        dimensions.append(
            (
                title,
                max((block.series[0][1].size, *(len(values) for _, values, _ in _boundary_export_columns(block.result)))) + 1,
                identity_column_count + len(block.series) + len(_boundary_export_columns(block.result)),
            )
        )
    if lineage_row_count:
        dimensions.append(
            (
                _worksheet_title(
                    "_Kindred Lineage",
                    used,
                ),
                lineage_row_count + 1,
                8,
            )
        )
    if metadata_row_count:
        dimensions.append(
            (
                _worksheet_title(
                    "_Kindred Result Metadata",
                    used,
                ),
                metadata_row_count + 1,
                3,
            )
        )
    return tuple(dimensions)


def _observation_provenance(
    source: ExperimentalLayerInput,
) -> Mapping[str, object]:
    payload: dict[str, object] = {
        "dataset_id": source.axis.dataset_id,
        "raw_observation_id": source.axis.raw_observation_id,
        "dataset_name": source.dataset_name,
        "raw_column": source.raw_column,
    }
    if source.model_target is not None:
        payload["model_target"] = source.model_target
    if source.set_id is not None:
        payload["set_id"] = source.set_id
    return MappingProxyType(payload)


def _same_source_row_values(
    x_source: ExperimentalLayerInput,
    y_source: ExperimentalLayerInput,
) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray | None] | None:
    """Pair experimental X/Y by the admitted dataset row identities."""

    if x_source.axis.dataset_id != y_source.axis.dataset_id:
        return None
    x_by_row = dict(zip(x_source.source_rows, x_source.y, strict=True))
    x_values: list[float] = []
    y_values: list[float] = []
    paired_rows: list[int] = []
    paired_times: list[float] = []
    for index, (row, value) in enumerate(zip(y_source.source_rows, y_source.y, strict=True)):
        if row not in x_by_row:
            continue
        x_values.append(float(x_by_row[row]))
        y_values.append(float(value))
        paired_rows.append(int(row))
        if y_source.t:
            paired_times.append(y_source.t[index])
    if not x_values:
        return None
    paired_x = np.asarray(x_values, dtype=float)
    paired_y = np.asarray(y_values, dtype=float)
    paired_source_rows = np.asarray(paired_rows, dtype=np.int64)
    paired_x.setflags(write=False)
    paired_y.setflags(write=False)
    paired_source_rows.setflags(write=False)
    sample_time = np.asarray(paired_times, dtype=float) if y_source.t else None
    if sample_time is not None:
        sample_time.setflags(write=False)
    return paired_x, paired_y, paired_source_rows, sample_time


def _annotations_from_execution(
    result: SimulationResultValue,
) -> tuple[Mapping[str, object], ...]:
    return result.intervention_execution.annotation_projection(
        execution_start=float(result.t[0]), execution_end=float(result.t[-1]),
    )


def intervention_label(
    annotation: Mapping[str, object], *, volume_unit: Unit, time_unit: Unit,
) -> str:
    """Format retained action amounts without evaluating their declarations."""
    action = annotation["physical_action"]
    if action is None:
        return html.escape(str(annotation["label"]))

    def quantity(fields, name, source, display):
        parameter = fields.get(name + "_param")
        value = html.escape(str(parameter)) if parameter is not None else f"{convert_value(fields[name], source, display):g}"
        unit_html = format_unit_for_display(display)
        return f"{value} {unit_html}".strip()

    op = action["op"]
    if op == "scale":
        factor = action["factor"]
        text = f"Scale all concentrations ×{factor:g} (volume ÷{factor:g})"
    elif op == "withdraw":
        text = f"Withdraw aliquot {quantity(action, 'volume', CUBIC_METRE, volume_unit)}"
    else:
        solution = action["solution"]
        if op == "flux":
            amount = quantity(solution, "volume_rate", CUBIC_METRE.divide(SECOND),
                              volume_unit.divide(time_unit))
        else:
            amount = quantity(solution, "volume", CUBIC_METRE, volume_unit)
        stock = ", ".join(f"{html.escape(item['species'])}: {quantity(item, 'value', MOLAR, MOLAR)}"
                          for item in solution["stock"]) or "pure solvent"
        text = f"{'Solution inflow' if op == 'flux' else 'Add solution'} {amount} ({stock})"
    if "initial_volume" in action or "initial_volume_param" in action:
        text += f" (volume at simulation start {quantity(action, 'initial_volume', CUBIC_METRE, volume_unit)})"
    return html.escape(str(annotation.get("label_prefix", ""))) + text


def _owned_annotations(
    result: SimulationResultValue, role: str, set_id: str,
) -> tuple[_DisplayAnnotation, ...]:
    return tuple(
        (role, set_id, annotation, 0.42 if role == "ghost" else 1.0)
        for annotation in _annotations_from_execution(result)
    )


def compose_interactive_display(
    result: SimulationResultValue,
    *,
    tab_id: str,
    set_id: str,
    render_input: ResultRenderInput,
    prepared_dsl: PreparedDslRequestSet,
) -> InteractiveDisplayProjection | None:
    # Reuse the static composer's resolved X, plot membership and colours so
    # interactive Preview never re-derives them from whichever block exists.
    selected_x = render_input.simulation_x_series
    if isinstance(selected_x, ExperimentalAxisSelection):
        return None
    label = dict(render_input.plot_set_labels).get(set_id)
    if label is None:
        return None
    wanted = None if render_input.selected_y_uses_default_all else set(render_input.selected_y_series)
    name_colors = {
        selection.series_name: color
        for selection, color in render_input.color_by_axis_selection
        if isinstance(selection, SimulationSeriesAxisSelection)
    }
    live = _result_block(
        set_id=set_id, label=label, result=result, prepared_dsl=prepared_dsl,
        kind=ResultLayerKind.PREVIEW,
        color_by_axis_selection=dict(render_input.color_by_axis_selection),
    )
    reference = next((block for block in render_input.blocks if block.block_id == set_id
                      and block.kind is ResultLayerKind.REFERENCE), None)
    curves = []
    annotations = []
    sources = [(result, live, "foreground")]
    if reference is not None and render_input.reference_lines_visible:
        sources.append((reference.result, reference, "ghost"))
    for value, style, role in sources:
        series = _series_tuple(value)
        colors = name_colors
        semantic = replace(
            style, result=value, t=value.t, sample_time=value.t, first_index=0,
            series=series, series_axes=tuple((name, SimulationSeriesAxisSelection(name)) for name, _ in series),
            x_axis=SimulationSeriesAxisSelection("t"), x_name="t",
            quantity_units=value.quantity_units, presentation_units=value.presentation_units,
            segment_ids=value.segment_ids, segment_keys=value.segment_keys,
            discontinuity_ids=value.discontinuity_ids, presentation_boundaries=value.presentation_boundaries,
            color_by_series=tuple((name, colors.get(name, "#3b82f6")) for name, _ in series),
        )
        vessel = _volume_block(semantic)
        for source in (semantic,) if vessel is None else (semantic, vessel):
            block = _data_block(source, x_axis=selected_x, y_axes=wanted)
            if block is None:
                continue
            axes = dict(block.series_axes)
            for name, values in block.series:
                chunks = physical_curve_chunks(value, selected_x, ((name, axes[name]),), range(block.first_index, len(value.t)))
                curves.append(InteractiveDisplayCurve(
                    key=("simulation", role, set_id, axes[name]), label=f"{style.label}: {name}",
                    x=block.t, y=values, sample_time=block.sample_time,
                    color=dict(block.color_by_series)[name], source=block, dashed=role == "ghost",
                    opacity=.42 if role == "ghost" else 1.,
                    segment_ids=block.segment_ids, segment_keys=block.segment_keys,
                    discontinuity_ids=block.discontinuity_ids,
                    physical_chunks=chunks,
                ))
        if _is_time_selection(selected_x):
            annotations.extend(_owned_annotations(value, role, set_id))
    if not any(curve.key[1] == "foreground" for curve in curves):
        return None
    return InteractiveDisplayProjection(
        tab_id=tab_id, set_id=set_id, x_label=render_input.x_label,
        y_label=render_input.y_label, curves=tuple(curves), intervention_annotations=tuple(annotations),
    )


def compose_result_input(
    *,
    tab_id: str,
    current_set_id: str | None,
    visible_sets: Sequence[VisibleResultSet],
    repository: ResultRepository,
    reference_lines_visible: bool,
    intervention_annotations_visible: bool,
    intervention_labels_visible: bool,
    active_overlay_layers: Sequence[ExperimentalLayerInput] = (),
    color_by_axis_selection: Mapping[AxisSeriesSelection, str] = MappingProxyType({}),
    experimental_point_shape: str = "o",
    experimental_point_size: float = 6.0,
    experimental_point_opacity: float = 0.85,
    selected_x_series: AxisSeriesSelection = SimulationSeriesAxisSelection("t"),
    selected_y_series: Sequence[AxisSeriesSelection] = (),
    selected_y_uses_default_all: bool = True,
) -> ResultRenderInput:
    tab_id = _required_id(tab_id, "tab_id")
    if repository.tab_id != tab_id:
        raise ValueError("repository does not belong to requested tab.")
    current_id = str(current_set_id or "")
    preview_ids = set(repository.preview_set_ids())
    colors = {
        _axis_selection(selection, field="configured color axis"):
        _required_id(color, "configured axis color")
        for selection, color in color_by_axis_selection.items()
    }
    selected_x = _axis_selection(selected_x_series, field="selected_x_series")
    selected_y = tuple(_axis_selection(item, field="selected_y_series") for item in selected_y_series)
    if len(selected_y) != len(set(selected_y)):
        raise ValueError("selected_y_series must be unique.")
    point_size = float(experimental_point_size)
    point_opacity = float(experimental_point_opacity)
    if not np.isfinite(point_size) or point_size <= 0:
        raise ValueError("experimental_point_size must be positive.")
    if not np.isfinite(point_opacity) or not 0 <= point_opacity <= 1:
        raise ValueError("experimental_point_opacity must be between zero and one.")
    active = tuple(active_overlay_layers)
    experimental_by_axis = {item.axis: item for item in active}
    if len(experimental_by_axis) != len(active):
        raise ValueError("active overlay layers duplicate experimental axis identity.")
    selected_model_y = {
        item.series_name for item in selected_y
        if isinstance(item, SimulationSeriesAxisSelection)
    }
    mapped_overlay_locks = tuple(
        (item.axis, str(item.model_target))
        for item in active
        if item.model_target is not None
        and (bool(selected_y_uses_default_all) or item.model_target in selected_model_y)
        and item.axis not in selected_y
    )
    effective_y = tuple(
        dict.fromkeys(
            (
                *selected_y,
                *(axis for axis, _model_target in mapped_overlay_locks),
            )
        )
    )
    wanted_y = None if bool(selected_y_uses_default_all) else set(effective_y)
    selected_x_source: ExperimentalLayerInput | None = None
    if isinstance(selected_x, ExperimentalAxisSelection):
        selected_x_source = experimental_by_axis.get(selected_x)
    simulation_x = (
        SimulationSeriesAxisSelection(selected_x_source.model_target)
        if isinstance(selected_x, ExperimentalAxisSelection)
        and selected_x_source is not None
        and selected_x_source.model_target is not None
        else selected_x
    )

    blocks: list[ResultDataBlock] = []
    unavailable_labels: list[str] = []
    candidates: list[AxisSelectionCandidate] = []
    candidate_ids: set[AxisSeriesSelection] = set()

    def add_candidate(selection: AxisSeriesSelection, label: str) -> None:
        if selection not in candidate_ids:
            candidate_ids.add(selection)
            candidates.append(AxisSelectionCandidate(selection, label))

    add_candidate(SimulationSeriesAxisSelection("t"), "t")
    for selection in colors:
        if isinstance(selection, SimulationSeriesAxisSelection):
            add_candidate(selection, selection.series_name)
    annotations: list[_DisplayAnnotation] = []
    warning_facts: list[tuple[str, tuple[str, ...]]] = []
    warning_ids: set[int] = set()
    seen_sets: set[str] = set()

    def add_warning(label: str, result: SimulationResultValue, visible: bool) -> None:
        has_algebra_warnings = bool(
            result.failed_algebra_outputs
            or result.omitted_algebra_outputs
            or result.algebra_errors
        )
        if visible and has_algebra_warnings and id(result) not in warning_ids:
            warning_ids.add(id(result))
            warning_facts.append((label, tuple(_result_warning_details(result))))

    for visible in visible_sets:
        if not isinstance(visible, VisibleResultSet):
            raise TypeError("visible_sets must contain VisibleResultSet values.")
        if visible.set_id in seen_sets:
            raise ValueError(f"visible set {visible.set_id!r} is duplicated.")
        seen_sets.add(visible.set_id)
        foreground = repository.current_result(visible.set_id)
        foreground_kind = ResultLayerKind.PREVIEW if visible.set_id in preview_ids else ResultLayerKind.EXPLICIT
        semantic: ResultDataBlock | None = None
        physical: ResultDataBlock | None = None
        if foreground is not None:
            semantic = _result_block(
                set_id=visible.set_id, label=visible.label, result=foreground,
                prepared_dsl=repository.current_request(visible.set_id).prepared_dsl,
                kind=foreground_kind, color_by_axis_selection=colors,
            )
            for name, selection in semantic.series_axes:
                add_candidate(selection, name)
            physical = _data_block(
                semantic, x_axis=simulation_x, y_axes=wanted_y,
            )
            if physical is not None:
                blocks.append(physical)
            else:
                if _has_selected_y_series(
                    semantic,
                    y_axes=wanted_y,
                ):
                    unavailable_labels.append(semantic.label)
                blocks.append(replace(semantic, visible=False))
            vessel = _volume_block(semantic)
            if vessel is not None:
                add_candidate(VesselVolumeAxisSelection(), "Vessel volume")
                vessel = _data_block(vessel, x_axis=simulation_x, y_axes=wanted_y)
                if vessel is not None:
                    blocks.append(vessel)
            add_warning(semantic.label, foreground, True)
            annotations.extend(_owned_annotations(foreground, "foreground", visible.set_id))
        reference = repository.reference_result(visible.set_id)
        reference_semantic: ResultDataBlock | None = None
        if reference is not None:
            reference_semantic = _result_block(
                set_id=visible.set_id, label=visible.label, result=reference,
                prepared_dsl=repository.reference_request(visible.set_id).prepared_dsl,
                kind=ResultLayerKind.REFERENCE,
                color_by_axis_selection=colors, visible=bool(reference_lines_visible),
            )
            for name, selection in reference_semantic.series_axes:
                add_candidate(selection, name)
            vessel = _volume_block(reference_semantic)
            if vessel is not None:
                add_candidate(VesselVolumeAxisSelection(), "Vessel volume")
                vessel = _data_block(vessel, x_axis=simulation_x, y_axes=wanted_y)
                if vessel is not None:
                    blocks.append(vessel)
            if reference_lines_visible:
                reference_physical = _data_block(
                    reference_semantic, x_axis=simulation_x, y_axes=wanted_y,
                )
                if reference_physical is not None:
                    blocks.append(reference_physical)
                else:
                    if _has_selected_y_series(
                        reference_semantic,
                        y_axes=wanted_y,
                    ):
                        unavailable_labels.append(reference_semantic.label)
                    blocks.append(replace(reference_semantic, visible=False))
                add_warning(reference_semantic.label, reference, True)
                annotations.extend(_owned_annotations(reference, "ghost", visible.set_id))
            else:
                blocks.append(reference_semantic)
    analysis_block = _select_analysis_block(blocks, current_id)
    statistics_names = (
        _statistics_series_names(analysis_block, y_axes=wanted_y)
        if analysis_block is not None
        else ()
    )
    for experimental in active:
        y_axis = experimental.axis
        label = f"{experimental.dataset_name}: {experimental.raw_column}"
        if experimental.model_target:
            label = f"{label} → {experimental.model_target}"
        add_candidate(y_axis, label)
        x_values = np.array(experimental.t, dtype=float, copy=True)
        y_values = np.array(experimental.y, dtype=float, copy=True)
        x_name = "t"
        x_unit = SECOND
        x_axis: AxisSeriesSelection = SimulationSeriesAxisSelection("t")
        x_provenance: Mapping[str, object] = MappingProxyType({})
        source_rows = np.array(experimental.source_rows, dtype=np.int64, copy=True)
        x_values.setflags(write=False)
        y_values.setflags(write=False)
        source_rows.setflags(write=False)
        sample_time = x_values if experimental.t else None
        unavailable = _is_time_selection(selected_x) and not experimental.t
        if not _is_time_selection(selected_x):
            x_source = (
                next(
                    (
                        item for item in active
                        if isinstance(selected_x, SimulationSeriesAxisSelection)
                        and item.axis.dataset_id == experimental.axis.dataset_id
                        and item.model_target == selected_x.series_name
                    ),
                    None,
                )
                if isinstance(selected_x, SimulationSeriesAxisSelection)
                else selected_x_source
            )
            if (
                x_source is None
                or x_source.axis.dataset_id != experimental.axis.dataset_id
            ):
                unavailable = True
            else:
                paired = _same_source_row_values(x_source, experimental)
                if paired is None:
                    unavailable = True
                else:
                    x_values, y_values, source_rows, sample_time = paired
                    x_name = x_source.raw_column
                    x_unit = x_source.unit
                    x_axis = selected_x
                    x_provenance = _observation_provenance(x_source)
        if unavailable:
            if wanted_y is None or y_axis in wanted_y:
                unavailable_labels.append(label)
            continue
        if wanted_y is not None and y_axis not in wanted_y:
            continue
        blocks.append(
            ResultDataBlock(
                block_id=experimental.axis,
                label=label,
                kind=ResultLayerKind.EXPERIMENTAL,
                t=x_values,
                sample_time=sample_time,
                time_column_name=experimental.time_column_name,
                series=((experimental.raw_column, y_values),),
                series_axes=((experimental.raw_column, y_axis),),
                x_axis=x_axis,
                result=None,
                segment_ids=None,
                segment_keys=None,
                discontinuity_ids=None,
                source_row=source_rows,
                x_name=x_name,
                quantity_units=((x_name, x_unit), (experimental.raw_column, experimental.unit)),
                series_provenance=(
                    (experimental.raw_column, _observation_provenance(experimental)),
                ),
                x_provenance=x_provenance,
                color_by_series=((experimental.raw_column, colors[y_axis]),),
                opacity=point_opacity,
                point_shape=experimental_point_shape,
                markers=True,
                point_size=point_size,
            )
        )

    owned_annotations = tuple(annotations) if _is_time_selection(selected_x) else ()
    if selected_x not in candidate_ids:
        add_candidate(
            selected_x,
            selected_x.series_name if isinstance(selected_x, SimulationSeriesAxisSelection)
            else "Vessel volume" if isinstance(selected_x, VesselVolumeAxisSelection)
            else f"{selected_x.dataset_id}: {selected_x.raw_observation_id}",
        )
    selected_candidate = next(item for item in candidates if item.selection == selected_x)
    def axis_units(axis):
        selected = [(block, block.x_name) for block in blocks if block.visible] if axis == "x" else [
            (block, name) for block in blocks if block.visible for name, _values in block.series]
        quantities = [block.unit_for(name, x=axis == "x") for block, name in selected]
        if not quantities or any(unit is None for unit in quantities):
            return None, ()
        unit = quantities[0]
        if any(other != unit for other in quantities):
            return None, ()
        preferred = tuple(dict(block.presentation_units).get(name, unit) for block, name in selected)
        preferred = tuple(item for item in preferred if item is not None and item.compatible(unit))
        return unit, compatible_units(unit, preferred)

    x_unit, x_options = axis_units("x")
    y_unit, y_options = axis_units("y")

    return ResultRenderInput(
        tab_id=tab_id, current_set_id=current_id, blocks=tuple(blocks),
        statistics_series_names=statistics_names,
        selected_x_series=selected_x, selected_y_series=selected_y,
        mapped_overlay_locks=mapped_overlay_locks,
        selected_y_uses_default_all=bool(selected_y_uses_default_all),
        x_label="Time (s)" if _is_time_selection(selected_x) else selected_candidate.label,
        y_label="Concentration" if y_unit is not None and y_unit.compatible(MOLAR) else "Value",
        x_unit=x_unit, y_unit=y_unit,
        x_unit_options=x_options, y_unit_options=y_options,
        axis_candidates=tuple(candidates),
        reference_lines_visible=bool(reference_lines_visible),
        intervention_annotations=owned_annotations,
        intervention_annotations_visible=intervention_annotations_visible,
        intervention_labels_visible=intervention_labels_visible,
        unavailable_x_diagnostic=_selected_x_unavailable_diagnostic(selected_x, unavailable_labels),
        completion_warning_details=tuple(warning_facts),
        plot_set_labels=tuple((item.set_id, item.label) for item in visible_sets),
        color_by_axis_selection=tuple(colors.items()),
        simulation_x_series=simulation_x,
    )

def iter_blocks_tsv_chunks(
    blocks: Sequence[ResultDataBlock],
    *,
    rows_per_chunk: int = 2048,
) -> Iterator[str]:
    block_values = tuple(blocks)
    columns: list[tuple[str, Sequence[object], str]] = []
    previous_group = None
    for block in block_values:
        group = (
            block.kind,
            block.block_id.dataset_id
            if block.kind is ResultLayerKind.EXPERIMENTAL
            else block.block_id,
        )
        if columns and group != previous_group:
            columns.append(("", (), "text"))
        previous_group = group
        columns.append(
            (
                _export_column_header(
                    block,
                    block.x_name,
                    block.x_provenance, x=True,
                ),
                block.t,
                "float",
            )
        )
        if block.kind is ResultLayerKind.EXPERIMENTAL:
            if block.source_row is None:
                raise ValueError("experimental Copy block lacks source_row.")
            columns.append(
                (
                    encode_clipboard_text(f"{block.label} — row_id"),
                    block.source_row,
                    "int",
                )
            )
        else:
            if (
                block.segment_ids is None
                or block.segment_keys is None
                or block.discontinuity_ids is None
            ):
                raise ValueError("simulation Copy block lacks executor identity.")
            columns.extend(
                (
                    (
                        encode_clipboard_text(f"{block.label} — segment_id"),
                        block.segment_ids,
                        "int",
                    ),
                    (
                        encode_clipboard_text(f"{block.label} — segment_key"),
                        block.segment_keys,
                        "text",
                    ),
                    (
                        encode_clipboard_text(f"{block.label} — discontinuity_id"),
                        block.discontinuity_ids,
                        "int",
                    ),
                )
            )
        provenance_by_name = dict(block.series_provenance)
        columns.extend(
            (
                _export_column_header(
                    block,
                    name,
                    provenance_by_name.get(name, MappingProxyType({})),
                ),
                values,
                "float",
            )
            for name, values in block.series
        )
        columns.extend((encode_clipboard_text(f"{block.label} — {name}"), values, kind)
                       for name, values, kind in _boundary_export_columns(block.result))
    if not columns:
        raise ValueError("No retained result or qualifying mapped dataset data.")
    row_count = max(len(values) for _header, values, _kind in columns)
    cell_count = (row_count + 1) * len(columns)
    if cell_count > _MAX_COPY_CELLS:
        raise ValueError(
            "Copy payload exceeds the bounded full-resolution cell limit "
            f"({_MAX_COPY_CELLS:,})."
        )
    chunk_size = int(rows_per_chunk)
    if chunk_size < 1:
        raise ValueError("rows_per_chunk must be positive.")
    total_characters = 0
    header = encode_tsv_row(tuple(header for header, _values, _kind in columns))
    total_characters += len(header)
    if total_characters > _MAX_COPY_CHARACTERS:
        raise ValueError(
            "Copy payload exceeds the bounded output-size limit "
            f"({_MAX_COPY_CHARACTERS // (1024 * 1024)} MiB)."
        )
    yield header
    for start in range(0, row_count, chunk_size):
        output = io.StringIO()
        for index in range(start, min(start + chunk_size, row_count)):
            output.write(
                encode_tsv_row(
                    tuple(
                        (
                            repr(float(values[index]))
                            if kind == "float"
                            else (
                                str(int(values[index]))
                                if kind == "int"
                                else encode_clipboard_text(str(values[index]))
                            )
                        )
                        if index < len(values) and values[index] is not None
                        else ""
                        for _header, values, kind in columns
                    )
                )
            )
        chunk = output.getvalue()
        total_characters += len(chunk)
        if total_characters > _MAX_COPY_CHARACTERS:
            raise ValueError(
                "Copy payload exceeds the bounded output-size limit "
                f"({_MAX_COPY_CHARACTERS // (1024 * 1024)} MiB)."
            )
        yield chunk

def validate_xlsx_export(
    blocks: Sequence[ResultDataBlock],
) -> tuple[tuple[str, int, str], ...]:
    """Admit worksheet dimensions and prepare canonical metadata rows once."""

    block_values = tuple(blocks)
    if not block_values:
        raise ValueError("No retained result or qualifying mapped dataset data.")
    if any(not isinstance(block, ResultDataBlock) for block in block_values):
        raise TypeError(
            "XLSX export blocks must contain only ResultDataBlock values."
        )
    lineage_row_count = sum(
        int(bool(block.x_provenance))
        + sum(bool(provenance) for _, provenance in block.series_provenance)
        for block in block_values
    )
    metadata_rows = _xlsx_metadata_rows(block_values)
    dimensions = _planned_xlsx_worksheet_dimensions(
        block_values,
        lineage_row_count=lineage_row_count,
        metadata_row_count=len(metadata_rows),
    )
    _validate_xlsx_worksheet_dimensions(dimensions)
    return metadata_rows

def write_xlsx(
    blocks: Sequence[ResultDataBlock],
    destination: BinaryIO,
    *,
    cancellation_check: Callable[[], None] | None = None,
) -> XlsxExportCompletion:
    """Write full-resolution worksheets directly to the destination."""

    block_values = tuple(blocks)
    metadata_rows = validate_xlsx_export(block_values)
    output = _XlsxDestination(
        destination,
        cancellation_check=cancellation_check,
    )
    workbook = Workbook(write_only=True)
    used: set[str] = set()
    lineage_rows: list[tuple[str, str, Mapping[str, object]]] = []
    data_rows = 0
    archive: ZipFile | None = None
    original_failure: BaseException | None = None
    try:
        for block in block_values:
            if cancellation_check is not None:
                cancellation_check()
            title = _worksheet_title(block.label, used)
            sheet = workbook.create_sheet(title=title)
            if block.kind is ResultLayerKind.EXPERIMENTAL:
                if block.source_row is None:
                    raise ValueError("experimental XLSX block lacks source_row.")
                identity_headers = ["row_id"]
            else:
                if (
                    block.segment_ids is None
                    or block.segment_keys is None
                    or block.discontinuity_ids is None
                ):
                    raise ValueError(
                        "simulation XLSX block lacks executor identity."
                    )
                identity_headers = [
                    "segment_id",
                    "segment_key",
                    "discontinuity_id",
                ]
            volume_columns = _boundary_export_columns(block.result)
            _append_xlsx_row(
                sheet,
                (
                    block.unit_label(block.x_name, x=True),
                    *identity_headers,
                    *(block.unit_label(name) for name, _values in block.series),
                    *(name for name, _, _ in volume_columns),
                ),
            )
            row_count = max((block.series[0][1].size, *(len(values) for _, values, _ in volume_columns)))
            for index in range(row_count):
                if cancellation_check is not None and index % 256 == 0:
                    cancellation_check()
                if index >= block.series[0][1].size:
                    identity_values = [None] * len(identity_headers)
                elif block.kind is ResultLayerKind.EXPERIMENTAL:
                    identity_values: list[object] = [int(block.source_row[index])]
                else:
                    identity_values = [
                        int(block.segment_ids[index]),
                        block.segment_keys[index],
                        int(block.discontinuity_ids[index]),
                    ]
                _append_xlsx_row(
                    sheet,
                    (
                        float(block.t[index]) if index < block.t.size else None,
                        *identity_values,
                        *(float(values[index]) if index < len(values) else None for _name, values in block.series),
                        *((int(values[index]) if kind == "int" else float(values[index])) if index < len(values) and values[index] is not None else None for _, values, kind in volume_columns),
                    ),
                )
                data_rows += 1
            if block.x_provenance:
                lineage_rows.append((title, block.x_name, block.x_provenance))
            lineage_rows.extend(
                (title, name, provenance)
                for name, provenance in block.series_provenance
                if provenance
            )
        if lineage_rows:
            if cancellation_check is not None:
                cancellation_check()
            lineage_sheet = workbook.create_sheet(
                    title=_worksheet_title(
                    "_Kindred Lineage",
                    used,
                )
            )
            _append_xlsx_row(
                lineage_sheet,
                (
                    "worksheet",
                    "series",
                    "dataset_id",
                    "raw_observation_id",
                    "dataset_name",
                    "raw_column",
                    "model_target",
                    "set_id",
                ),
            )
            for title, name, provenance in lineage_rows:
                if cancellation_check is not None:
                    cancellation_check()
                _append_xlsx_row(
                    lineage_sheet,
                    (
                        title,
                        name,
                        str(provenance.get("dataset_id") or ""),
                        str(provenance.get("raw_observation_id") or ""),
                        str(provenance.get("dataset_name") or ""),
                        str(provenance.get("raw_column") or ""),
                        str(provenance.get("model_target") or ""),
                        str(provenance.get("set_id") or ""),
                    ),
                )
            lineage_sheet.sheet_state = "hidden"
        if metadata_rows:
            if cancellation_check is not None:
                cancellation_check()
            metadata_sheet = workbook.create_sheet(
                    title=_worksheet_title(
                    "_Kindred Result Metadata",
                    used,
                )
            )
            _append_xlsx_row(
                metadata_sheet,
                ("worksheet", "part", "repository_entry"),
            )
            for row in metadata_rows:
                if cancellation_check is not None:
                    cancellation_check()
                _append_xlsx_row(
                    metadata_sheet,
                    row,
                )
            metadata_sheet.sheet_state = "hidden"
        if cancellation_check is not None:
            cancellation_check()
        archive = ZipFile(
            output,
            "w",
            ZIP_DEFLATED,
            allowZip64=True,
        )
        ExcelWriter(workbook, archive).save()
        return XlsxExportCompletion(
            worksheet_count=len(used),
            data_row_count=data_rows,
            bytes_written=output.bytes_written,
        )
    except BaseException as exc:
        original_failure = exc
        raise
    finally:
        cleanup_failures: list[BaseException] = []
        if archive is not None:
            try:
                archive.close()
            except BaseException as exc:
                cleanup_failures.append(exc)
        cleanup_failures.extend(_cleanup_write_only_workbook(workbook))
        if original_failure is None and cleanup_failures:
            raise cleanup_failures[0]

def _worksheet_title(label: str, used: set[str]) -> str:
    invalid = set("[]:*?/\\")
    base = "".join(
        "_" if char in invalid or ord(char) < 32 else char for char in str(label)
    ).strip()
    base = base.strip("'") or "Result"
    base = base[:31]
    candidate = base
    suffix_index = 2
    while candidate.casefold() in used:
        suffix = f" ({suffix_index})"
        candidate = f"{base[: 31 - len(suffix)]}{suffix}"
        suffix_index += 1
    used.add(candidate.casefold())
    return candidate
